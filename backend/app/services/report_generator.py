"""
Report Generation Service

Generates sprint quality reports in PDF and Excel formats.
Aggregates data from quality metrics, test coverage, and defect tracking.
"""

import os
import uuid
from datetime import datetime, timezone
from typing import Optional, List, Dict
from io import BytesIO

from fpdf import FPDF
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Border, Side

from app.schemas.quality import (
    SprintQualityReport,
    ReportFormat,
    ReportSection,
    TestCoverageSection,
    PRQualitySection,
    ComponentHealthItem,
    RecommendationItem,
    QualitySummary,
)


class QualityReportPDF(FPDF):
    """Custom PDF class for quality reports."""

    def __init__(self, report: SprintQualityReport):
        super().__init__()
        self.report = report
        self.set_auto_page_break(auto=True, margin=15)

    def header(self):
        self.set_font("Helvetica", "B", 16)
        self.cell(0, 10, "Sprint Quality Report", border=0, ln=True, align="C")
        self.set_font("Helvetica", "", 10)
        project_name = self.report.project_name or f"Project #{self.report.project_id}"
        self.cell(0, 6, project_name, border=0, ln=True, align="C")
        if self.report.sprint_name:
            self.cell(0, 6, f"Sprint: {self.report.sprint_name}", border=0, ln=True, align="C")
        self.cell(
            0,
            6,
            f"Generated: {self.report.report_date.strftime('%Y-%m-%d %H:%M')}",
            border=0,
            ln=True,
            align="C",
        )
        self.ln(5)
        self.set_draw_color(200, 200, 200)
        self.line(10, self.get_y(), 200, self.get_y())
        self.ln(5)

    def footer(self):
        self.set_y(-15)
        self.set_font("Helvetica", "I", 8)
        self.set_text_color(128, 128, 128)
        self.cell(0, 10, f"Page {self.page_no()}/{{nb}}", align="C")

    def section_title(self, title: str):
        self.set_font("Helvetica", "B", 14)
        self.set_fill_color(240, 240, 240)
        self.cell(0, 10, title, border=0, ln=True, fill=True)
        self.ln(2)

    def subsection_title(self, title: str):
        self.set_font("Helvetica", "B", 11)
        self.cell(0, 8, title, border=0, ln=True)
        self.ln(1)

    def add_metric_row(self, label: str, value: str, color: tuple = (0, 0, 0)):
        self.set_font("Helvetica", "", 10)
        self.cell(70, 6, label, border=0)
        self.set_text_color(*color)
        self.set_font("Helvetica", "B", 10)
        self.cell(0, 6, value, border=0, ln=True)
        self.set_text_color(0, 0, 0)

    def add_quality_grade_box(self, grade: str, score: float):
        grade_colors = {
            "A": (46, 204, 113),
            "B": (52, 152, 219),
            "C": (241, 196, 15),
            "D": (230, 126, 34),
            "F": (231, 76, 60),
        }
        color = grade_colors.get(grade, (128, 128, 128))
        self.set_fill_color(*color)
        self.set_text_color(255, 255, 255)
        self.set_font("Helvetica", "B", 24)
        x_start = self.get_x()
        self.cell(40, 30, grade, border=0, fill=True, align="C")
        self.set_xy(x_start + 45, self.get_y() + 8)
        self.set_text_color(0, 0, 0)
        self.set_font("Helvetica", "", 12)
        self.cell(0, 6, f"Quality Score: {score:.1f}/100", ln=True)
        self.ln(20)

    def add_highlights_concerns(self, highlights: List[str], concerns: List[str]):
        if highlights:
            self.subsection_title("Key Highlights")
            self.set_font("Helvetica", "", 10)
            self.set_text_color(46, 204, 113)
            for h in highlights[:5]:
                self.cell(5, 6, "+")
                self.set_text_color(0, 0, 0)
                self.cell(0, 6, h, ln=True)
                self.set_text_color(46, 204, 113)
            self.set_text_color(0, 0, 0)
            self.ln(3)

        if concerns:
            self.subsection_title("Key Concerns")
            self.set_font("Helvetica", "", 10)
            self.set_text_color(231, 76, 60)
            for c in concerns[:5]:
                self.cell(5, 6, "!")
                self.set_text_color(0, 0, 0)
                self.cell(0, 6, c, ln=True)
                self.set_text_color(231, 76, 60)
            self.set_text_color(0, 0, 0)
            self.ln(3)


class ReportGeneratorService:
    """Service for generating quality reports."""

    def __init__(self, reports_dir: str = "/tmp/reports"):
        self.reports_dir = reports_dir
        os.makedirs(reports_dir, exist_ok=True)

    def generate_report_data(
        self,
        project_id: int,
        project_name: Optional[str] = None,
        sprint_id: Optional[int] = None,
        sprint_name: Optional[str] = None,
        quality_summary: Optional[QualitySummary] = None,
        test_coverage: Optional[TestCoverageSection] = None,
        pr_quality: Optional[PRQualitySection] = None,
        escaped_defects: Optional[List[Dict]] = None,
        component_health: Optional[List[ComponentHealthItem]] = None,
        coverage_trend: Optional[List[Dict]] = None,
        defect_trend: Optional[List[Dict]] = None,
    ) -> SprintQualityReport:
        """Generate the report data structure with computed metrics."""

        # Compute quality score and grade
        score = self._compute_quality_score(
            quality_summary=quality_summary,
            test_coverage=test_coverage,
            pr_quality=pr_quality,
        )
        grade = self._score_to_grade(score)

        # Generate highlights and concerns
        highlights, concerns = self._generate_highlights_concerns(
            quality_summary=quality_summary,
            test_coverage=test_coverage,
            pr_quality=pr_quality,
        )

        # Generate recommendations
        recommendations = self._generate_recommendations(
            quality_summary=quality_summary,
            test_coverage=test_coverage,
            pr_quality=pr_quality,
            component_health=component_health,
        )

        return SprintQualityReport(
            report_id=str(uuid.uuid4()),
            project_id=project_id,
            project_name=project_name,
            sprint_id=sprint_id,
            sprint_name=sprint_name,
            report_date=datetime.now(timezone.utc),
            overall_quality_score=score,
            quality_grade=grade,
            key_highlights=highlights,
            key_concerns=concerns,
            quality_summary=quality_summary,
            test_coverage=test_coverage,
            pr_quality=pr_quality,
            escaped_defects_summary={
                "total": quality_summary.total_escaped_defects if quality_summary else 0,
                "open": quality_summary.open_defects if quality_summary else 0,
                "resolved": quality_summary.resolved_defects if quality_summary else 0,
            }
            if quality_summary
            else None,
            recent_escaped_defects=[],
            component_health=component_health or [],
            coverage_trend=coverage_trend or [],
            defect_trend=defect_trend or [],
            recommendations=recommendations,
        )

    def _compute_quality_score(
        self,
        quality_summary: Optional[QualitySummary],
        test_coverage: Optional[TestCoverageSection],
        pr_quality: Optional[PRQualitySection],
    ) -> float:
        """Compute overall quality score 0-100."""
        scores = []
        weights = []

        # Coverage score (30% weight)
        if test_coverage and test_coverage.line_coverage is not None:
            coverage_score = min(100, test_coverage.line_coverage * 100 / 0.8)
            scores.append(coverage_score)
            weights.append(30)

        # Test pass rate (20% weight)
        if test_coverage and test_coverage.test_pass_rate:
            scores.append(test_coverage.test_pass_rate * 100)
            weights.append(20)

        # DRE score (25% weight)
        if quality_summary and quality_summary.defect_removal_efficiency is not None:
            scores.append(quality_summary.defect_removal_efficiency * 100)
            weights.append(25)

        # PR gate pass rate (15% weight)
        if pr_quality and pr_quality.gate_pass_rate:
            scores.append(pr_quality.gate_pass_rate * 100)
            weights.append(15)

        # Low escaped defects (10% weight)
        if quality_summary:
            escaped = quality_summary.total_escaped_defects or 0
            defect_score = max(0, 100 - escaped * 10)
            scores.append(defect_score)
            weights.append(10)

        if not scores:
            return 50.0

        weighted_sum = sum(s * w for s, w in zip(scores, weights))
        total_weight = sum(weights)
        return round(weighted_sum / total_weight, 1)

    def _score_to_grade(self, score: float) -> str:
        """Convert score to letter grade."""
        if score >= 90:
            return "A"
        elif score >= 80:
            return "B"
        elif score >= 70:
            return "C"
        elif score >= 60:
            return "D"
        return "F"

    def _generate_highlights_concerns(
        self,
        quality_summary: Optional[QualitySummary],
        test_coverage: Optional[TestCoverageSection],
        pr_quality: Optional[PRQualitySection],
    ) -> tuple:
        """Generate key highlights and concerns."""
        highlights = []
        concerns = []

        if test_coverage:
            if test_coverage.coverage_met:
                highlights.append(f"Test coverage target met: {test_coverage.line_coverage:.1%}")
            elif test_coverage.line_coverage is not None:
                if test_coverage.line_coverage < 0.5:
                    concerns.append(f"Low test coverage: {test_coverage.line_coverage:.1%}")

            if test_coverage.test_pass_rate >= 0.95:
                highlights.append(f"Excellent test pass rate: {test_coverage.test_pass_rate:.1%}")
            elif test_coverage.test_pass_rate < 0.9:
                concerns.append(f"Test pass rate below 90%: {test_coverage.test_pass_rate:.1%}")

            if test_coverage.flaky_test_count == 0:
                highlights.append("No flaky tests detected")
            elif test_coverage.flaky_test_count > 5:
                concerns.append(f"{test_coverage.flaky_test_count} flaky tests need attention")

        if quality_summary:
            if (
                quality_summary.defect_removal_efficiency
                and quality_summary.defect_removal_efficiency >= 0.95
            ):
                highlights.append(f"High DRE: {quality_summary.defect_removal_efficiency:.1%}")
            elif (
                quality_summary.defect_removal_efficiency
                and quality_summary.defect_removal_efficiency < 0.85
            ):
                concerns.append(
                    f"DRE below target: {quality_summary.defect_removal_efficiency:.1%}"
                )

            if quality_summary.critical_count > 0:
                concerns.append(
                    f"{quality_summary.critical_count} critical defects escaped to production"
                )

            if quality_summary.total_escaped_defects == 0:
                highlights.append("Zero escaped defects this period")

        if pr_quality:
            if pr_quality.gate_pass_rate >= 0.9:
                highlights.append(f"PR quality gates: {pr_quality.gate_pass_rate:.1%} pass rate")
            elif pr_quality.gate_pass_rate < 0.7:
                concerns.append(
                    f"PR gate pass rate needs improvement: {pr_quality.gate_pass_rate:.1%}"
                )

        return highlights[:5], concerns[:5]

    def _generate_recommendations(
        self,
        quality_summary: Optional[QualitySummary],
        test_coverage: Optional[TestCoverageSection],
        pr_quality: Optional[PRQualitySection],
        component_health: Optional[List[ComponentHealthItem]],
    ) -> List[RecommendationItem]:
        """Generate actionable recommendations."""
        recommendations = []

        if test_coverage:
            if test_coverage.line_coverage is not None and test_coverage.line_coverage < 0.8:
                recommendations.append(
                    RecommendationItem(
                        category="coverage",
                        priority="high",
                        recommendation=f"Increase test coverage from {test_coverage.line_coverage:.1%} to 80% target",
                        impact="Reduce escaped defects and improve confidence in releases",
                        effort="medium",
                    )
                )

            if test_coverage.flaky_test_count > 0:
                recommendations.append(
                    RecommendationItem(
                        category="testing",
                        priority="high" if test_coverage.flaky_test_count > 5 else "medium",
                        recommendation=f"Fix {test_coverage.flaky_test_count} flaky tests to improve CI reliability",
                        impact="Reduce false failures and improve developer productivity",
                        effort="medium",
                    )
                )

        if quality_summary:
            if (
                quality_summary.defect_removal_efficiency
                and quality_summary.defect_removal_efficiency < 0.9
            ):
                recommendations.append(
                    RecommendationItem(
                        category="defects",
                        priority="high",
                        recommendation="Improve defect detection by adding more unit and integration tests",
                        impact="Catch more defects before they escape to production",
                        effort="high",
                    )
                )

        if component_health:
            critical_components = [c for c in component_health if c.priority == "critical"]
            if critical_components:
                recommendations.append(
                    RecommendationItem(
                        category="components",
                        priority="critical",
                        recommendation=f"Address {len(critical_components)} components with critical risk scores",
                        impact="Reduce overall system risk and improve reliability",
                        effort="high",
                    )
                )

        if pr_quality and pr_quality.gate_pass_rate < 0.8:
            recommendations.append(
                RecommendationItem(
                    category="process",
                    priority="medium",
                    recommendation="Review and adjust PR quality gate criteria or improve pre-submission checks",
                    impact="Improve first-time PR approval rate",
                    effort="low",
                )
            )

        return recommendations[:10]

    def generate_pdf(self, report: SprintQualityReport, sections: List[ReportSection]) -> bytes:
        """Generate PDF report."""
        pdf = QualityReportPDF(report)
        pdf.alias_nb_pages()
        pdf.add_page()

        # Executive Summary
        if ReportSection.executive_summary in sections:
            pdf.section_title("Executive Summary")
            pdf.add_quality_grade_box(report.quality_grade, report.overall_quality_score)
            pdf.add_highlights_concerns(report.key_highlights, report.key_concerns)

        # Quality Metrics
        if ReportSection.quality_metrics in sections and report.quality_summary:
            pdf.section_title("Quality Metrics")
            qs = report.quality_summary
            pdf.add_metric_row("Total Escaped Defects", str(qs.total_escaped_defects))
            pdf.add_metric_row("Open Defects", str(qs.open_defects))
            pdf.add_metric_row("Resolved Defects", str(qs.resolved_defects))
            if qs.defect_removal_efficiency is not None:
                color = (46, 204, 113) if qs.defect_removal_efficiency >= 0.9 else (231, 76, 60)
                pdf.add_metric_row(
                    "Defect Removal Efficiency", f"{qs.defect_removal_efficiency:.1%}", color
                )
            if qs.defect_density is not None:
                pdf.add_metric_row("Defect Density", f"{qs.defect_density:.2f} per KLOC")
            if qs.mttr_hours is not None:
                pdf.add_metric_row("Mean Time to Resolve", f"{qs.mttr_hours:.1f} hours")
            pdf.ln(5)

        # Test Coverage
        if ReportSection.test_coverage in sections and report.test_coverage:
            pdf.section_title("Test Coverage")
            tc = report.test_coverage
            if tc.line_coverage is not None:
                color = (46, 204, 113) if tc.coverage_met else (231, 76, 60)
                pdf.add_metric_row("Line Coverage", f"{tc.line_coverage:.1%}", color)
            if tc.branch_coverage is not None:
                pdf.add_metric_row("Branch Coverage", f"{tc.branch_coverage:.1%}")
            pdf.add_metric_row("Coverage Target", f"{tc.coverage_target:.0%}")
            pdf.add_metric_row("Total Tests", str(tc.total_tests))
            pdf.add_metric_row("Passed Tests", str(tc.passed_tests))
            pdf.add_metric_row("Failed Tests", str(tc.failed_tests))
            pdf.add_metric_row("Test Pass Rate", f"{tc.test_pass_rate:.1%}")
            pdf.add_metric_row("Flaky Tests", str(tc.flaky_test_count))
            pdf.ln(5)

        # PR Quality
        if ReportSection.pr_quality in sections and report.pr_quality:
            pdf.section_title("PR Quality Gates")
            pr = report.pr_quality
            pdf.add_metric_row("Total PRs Merged", str(pr.total_prs_merged))
            pdf.add_metric_row("PRs Meeting Gates", str(pr.prs_meeting_gates))
            color = (46, 204, 113) if pr.gate_pass_rate >= 0.8 else (231, 76, 60)
            pdf.add_metric_row("Gate Pass Rate", f"{pr.gate_pass_rate:.1%}", color)
            if pr.avg_review_time_hours is not None:
                pdf.add_metric_row("Avg Review Time", f"{pr.avg_review_time_hours:.1f} hours")
            pdf.ln(5)

        # Component Health
        if ReportSection.component_health in sections and report.component_health:
            pdf.section_title("Component Health")
            pdf.set_font("Helvetica", "B", 9)
            pdf.cell(60, 6, "Component", border=1)
            pdf.cell(25, 6, "Coverage", border=1, align="C")
            pdf.cell(25, 6, "Defects", border=1, align="C")
            pdf.cell(25, 6, "Risk", border=1, align="C")
            pdf.cell(30, 6, "Priority", border=1, align="C", ln=True)
            pdf.set_font("Helvetica", "", 9)
            for comp in report.component_health[:15]:
                pdf.cell(60, 6, comp.component[:30], border=1)
                cov_str = f"{comp.coverage:.1%}" if comp.coverage else "N/A"
                pdf.cell(25, 6, cov_str, border=1, align="C")
                pdf.cell(25, 6, str(comp.defect_count), border=1, align="C")
                pdf.cell(25, 6, f"{comp.risk_score:.1f}", border=1, align="C")
                pdf.cell(30, 6, comp.priority.upper(), border=1, align="C", ln=True)
            pdf.ln(5)

        # Recommendations
        if ReportSection.recommendations in sections and report.recommendations:
            pdf.section_title("Recommendations")
            for i, rec in enumerate(report.recommendations[:10], 1):
                pdf.set_font("Helvetica", "B", 10)
                priority_colors = {
                    "critical": (231, 76, 60),
                    "high": (230, 126, 34),
                    "medium": (241, 196, 15),
                    "low": (46, 204, 113),
                }
                color = priority_colors.get(rec.priority, (128, 128, 128))
                pdf.set_text_color(*color)
                pdf.cell(0, 6, f"{i}. [{rec.priority.upper()}] {rec.category.upper()}", ln=True)
                pdf.set_text_color(0, 0, 0)
                pdf.set_font("Helvetica", "", 10)
                pdf.multi_cell(0, 5, rec.recommendation)
                pdf.set_font("Helvetica", "I", 9)
                pdf.cell(0, 5, f"Impact: {rec.impact} | Effort: {rec.effort}", ln=True)
                pdf.ln(2)

        return bytes(pdf.output())

    def generate_excel(self, report: SprintQualityReport, sections: List[ReportSection]) -> bytes:
        """Generate Excel report."""
        wb = Workbook()

        # Styles
        header_font = Font(bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
        thin_border = Border(
            left=Side(style="thin"),
            right=Side(style="thin"),
            top=Side(style="thin"),
            bottom=Side(style="thin"),
        )

        # Summary Sheet
        ws = wb.active
        ws.title = "Summary"
        ws["A1"] = "Sprint Quality Report"
        ws["A1"].font = Font(bold=True, size=16)
        ws["A2"] = f"Project: {report.project_name or report.project_id}"
        ws["A3"] = f"Sprint: {report.sprint_name or 'N/A'}"
        ws["A4"] = f"Generated: {report.report_date.strftime('%Y-%m-%d %H:%M')}"
        ws["A6"] = "Quality Grade"
        ws["B6"] = report.quality_grade
        ws["B6"].font = Font(bold=True, size=24)
        ws["A7"] = "Quality Score"
        ws["B7"] = f"{report.overall_quality_score}/100"

        # Key Metrics
        ws["A9"] = "Key Metrics"
        ws["A9"].font = Font(bold=True)
        row = 10
        if report.quality_summary:
            qs = report.quality_summary
            metrics = [
                ("Total Escaped Defects", qs.total_escaped_defects),
                ("Open Defects", qs.open_defects),
                ("Resolved Defects", qs.resolved_defects),
                (
                    "DRE",
                    f"{qs.defect_removal_efficiency:.1%}"
                    if qs.defect_removal_efficiency
                    else "N/A",
                ),
                ("Defect Density", f"{qs.defect_density:.2f}" if qs.defect_density else "N/A"),
            ]
            for label, value in metrics:
                ws[f"A{row}"] = label
                ws[f"B{row}"] = value
                row += 1

        if report.test_coverage:
            tc = report.test_coverage
            row += 1
            metrics = [
                ("Line Coverage", f"{tc.line_coverage:.1%}" if tc.line_coverage else "N/A"),
                ("Branch Coverage", f"{tc.branch_coverage:.1%}" if tc.branch_coverage else "N/A"),
                ("Test Pass Rate", f"{tc.test_pass_rate:.1%}"),
                ("Flaky Tests", tc.flaky_test_count),
            ]
            for label, value in metrics:
                ws[f"A{row}"] = label
                ws[f"B{row}"] = value
                row += 1

        # Component Health Sheet
        if ReportSection.component_health in sections and report.component_health:
            ws_comp = wb.create_sheet("Component Health")
            headers = ["Component", "Coverage", "Defects", "Flaky Tests", "Risk Score", "Priority"]
            for col, header in enumerate(headers, 1):
                cell = ws_comp.cell(row=1, column=col, value=header)
                cell.font = header_font
                cell.fill = header_fill
                cell.border = thin_border

            for row_idx, comp in enumerate(report.component_health, 2):
                ws_comp.cell(row=row_idx, column=1, value=comp.component).border = thin_border
                ws_comp.cell(
                    row=row_idx, column=2, value=f"{comp.coverage:.1%}" if comp.coverage else "N/A"
                ).border = thin_border
                ws_comp.cell(row=row_idx, column=3, value=comp.defect_count).border = thin_border
                ws_comp.cell(
                    row=row_idx, column=4, value=comp.flaky_test_count
                ).border = thin_border
                ws_comp.cell(
                    row=row_idx, column=5, value=round(comp.risk_score, 2)
                ).border = thin_border
                ws_comp.cell(
                    row=row_idx, column=6, value=comp.priority.upper()
                ).border = thin_border

            # Auto-width columns
            for col_cells in ws_comp.iter_cols():
                if not col_cells:
                    continue
                max_length = max(len(str(cell.value or "")) for cell in col_cells)
                first_cell = col_cells[0]
                if hasattr(first_cell, "column_letter"):
                    ws_comp.column_dimensions[first_cell.column_letter].width = min(
                        50, max_length + 2
                    )

        # Recommendations Sheet
        if ReportSection.recommendations in sections and report.recommendations:
            ws_rec = wb.create_sheet("Recommendations")
            headers = ["Priority", "Category", "Recommendation", "Impact", "Effort"]
            for col, header in enumerate(headers, 1):
                cell = ws_rec.cell(row=1, column=col, value=header)
                cell.font = header_font
                cell.fill = header_fill
                cell.border = thin_border

            for row_idx, rec in enumerate(report.recommendations, 2):
                ws_rec.cell(row=row_idx, column=1, value=rec.priority.upper()).border = thin_border
                ws_rec.cell(row=row_idx, column=2, value=rec.category).border = thin_border
                ws_rec.cell(row=row_idx, column=3, value=rec.recommendation).border = thin_border
                ws_rec.cell(row=row_idx, column=4, value=rec.impact).border = thin_border
                ws_rec.cell(row=row_idx, column=5, value=rec.effort).border = thin_border

            for col_cells in ws_rec.iter_cols():
                if not col_cells:
                    continue
                max_length = max(len(str(cell.value or "")) for cell in col_cells)
                first_cell = col_cells[0]
                if hasattr(first_cell, "column_letter"):
                    ws_rec.column_dimensions[first_cell.column_letter].width = min(
                        60, max_length + 2
                    )

        # Save to bytes
        output = BytesIO()
        wb.save(output)
        output.seek(0)
        return output.getvalue()

    def save_report(
        self,
        report: SprintQualityReport,
        format: ReportFormat,
        sections: List[ReportSection],
    ) -> tuple:
        """Generate and save report to disk. Returns (filename, filepath, size_bytes)."""
        timestamp = datetime.now(timezone.utc).strftime("%Y%m%d_%H%M%S")
        project_slug = (
            (report.project_name or f"project_{report.project_id}").replace(" ", "_").lower()
        )

        if format == ReportFormat.pdf:
            content = self.generate_pdf(report, sections)
            filename = f"quality_report_{project_slug}_{timestamp}.pdf"
        elif format == ReportFormat.excel:
            content = self.generate_excel(report, sections)
            filename = f"quality_report_{project_slug}_{timestamp}.xlsx"
        elif format == ReportFormat.json:
            content = report.model_dump_json(indent=2).encode()
            filename = f"quality_report_{project_slug}_{timestamp}.json"
        else:
            content = report.model_dump_json(indent=2).encode()
            filename = f"quality_report_{project_slug}_{timestamp}.json"

        filepath = os.path.join(self.reports_dir, filename)
        with open(filepath, "wb") as f:
            f.write(content)

        return filename, filepath, len(content)
